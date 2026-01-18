# rag/loader.py
from typing import List
from pathlib import Path

from langchain.docstore.document import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from docx import Document as DocxDocument
from langchain.text_splitter import CharacterTextSplitter

# =========================
# 配置
# =========================



# =========================
# 对外唯一接口
# =========================
def load_and_split_docs(files):
    all_docs = []
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=300,
        chunk_overlap=50,
        separators=["\n\n", "\n"],
        length_function=len     # 使用字符数作为长度
    )

    for file in files:
        doc = DocxDocument(str(file))
        # 去掉空段落，并拼接成连续文本
        full_text = "".join([para.text.strip() for para in doc.paragraphs if para.text.strip()])
        doc_obj = Document(page_content=full_text, metadata={"source": Path(file).name})
        split_docs = splitter.split_documents([doc_obj])
        all_docs.extend(split_docs)

    return all_docs